"""Fase 7: relatório (PDF + planilha), fila, redação conferida, comparativo de renda fixa e calculadoras (à mão)."""

import json
from datetime import date

import pymupdf
import pytest
from openpyxl import load_workbook

from quiron.servicos import calculadoras as calc
from quiron.servicos.analise import redacao
from quiron.servicos.analise.fila import Fila, TIPOS, tipo
from quiron.servicos.analise.relatorio import Grafico, Relatorio, Secao, Tabela, formatar
from quiron.servicos.analise.tipos import renda_fixa as rf


def _relatorio() -> Relatorio:
    return Relatorio(
        "Teste — relatório", "teste", subtitulo="sub", resumo=["Conclusão com **R$ 110.000,00**."],
        premissas=["Taxa de 10% a.a."], fontes=["📊 Fonte X — 04/10/2026"], limitacoes=["Exemplo."],
        fatos={"liquido": 110000.0, "taxa": 10.0},
        secoes=[Secao("Resultado", "Texto **forte**.",
                      [Tabela("Tabela", ["Opção", "Valor"], [["A", 110000.0], ["B", 105000.0]], ["texto", "brl"], "nota")],
                      [Grafico("Barras", "barras_h", ["A", "B"], {"v": [110000, 105000]}, "brl"),
                       Grafico("Linhas", "linhas", ["1", "2", "3"], {"A": [1, 2, 3], "B": [1, 1.5, 2]}, "num", "meses")])])


def test_relatorio_gera_pdf_planilha_e_markdown(tmp_path):
    caminhos = _relatorio().salvar(tmp_path / "r")
    doc = pymupdf.open(caminhos["pdf"])
    texto = "".join(p.get_text() for p in doc)
    assert "Teste — relatório" in texto and "R$ 110.000,00" in texto and "Premissas" in texto and "Limitações" in texto
    assert "não constitui recomendação" in doc[-1].get_text()  # rodapé de compliance em todas as páginas
    assert sum(len(p.get_images()) for p in doc) == 2
    alturas = [i["bbox"][3] - i["bbox"][1] for p in doc for i in p.get_image_info()]
    assert min(alturas) > 100  # gráfico nunca encolhido
    wb = load_workbook(caminhos["planilha"])
    assert wb.sheetnames[:2] == ["Resumo", "Tabela"] and wb["Tabela"]["B2"].value == 110000.0
    assert wb["Tabela"]["B2"].number_format.startswith('"R$"')
    md = caminhos["markdown"].read_text(encoding="utf-8")
    assert "| A | R$ 110.000,00 |" in md
    volta = Relatorio.de_dict(json.loads(caminhos["json"].read_text(encoding="utf-8")))
    assert volta.secoes[0].tabelas[0].linhas[0][1] == 110000.0
    assert formatar(1234.5, "brl") == "R$ 1.234,50" and formatar(12.345, "pct") == "12,35%"


def test_conferencia_de_numeros_pega_invencao():
    rel = _relatorio()
    refs = redacao.referencias(rel)
    assert redacao.conferir("Rende R$ 110.000,00, a 10% ao ano, 5 mil a mais (R$ 5.000,00).", refs) == []
    assert redacao.conferir("O CDB rende 17,3% ao ano.", refs) == [17.3]


def test_redacao_usa_texto_do_modelo_ou_avisa(monkeypatch):
    rel = _relatorio()
    bom = json.dumps({"resumo": ["A vence com R$ 110.000,00."], "analise": "Porque rende 10% a.a."})
    redacao.redigir(rel, "contestar", mock_response=bom)
    assert rel.resumo == ["A vence com R$ 110.000,00."] and rel.secoes[0].titulo == "Advogado do diabo"
    rel2 = _relatorio()
    ruim = json.dumps({"resumo": ["A rende 23,7% a.a."], "analise": "x"})
    redacao.redigir(rel2, "entregar", mock_response=ruim)
    assert any("não conferidos" in a for a in rel2.avisos)


def test_fila_processa_arquiva_busca_e_entrega(tmp_path, monkeypatch):
    monkeypatch.setattr("quiron.servicos.analise.fila.carregar_tipos", lambda: TIPOS)

    @tipo("teste_ok", "teste")
    def _ok(params, modo):
        r = _relatorio()
        r.titulo = f"Estudo de debêntures {params['n']}"
        return r

    @tipo("teste_erro", "teste")
    def _erro(params, modo):
        raise RuntimeError("fonte fora do ar")

    f = Fila(tmp_path / "a.db", tmp_path / "rel")
    with pytest.raises(ValueError):
        f.pedir("nao_existe")
    t1 = f.pedir("teste_ok", {"n": 1}, "debater", origem="telegram")
    t2 = f.pedir("teste_erro", {}, origem="telegram")
    assert f.posicao(t2.id) == 1
    assert f.processar_uma().situacao == "pronta"
    assert f.processar_uma().erro.startswith("RuntimeError")
    assert f.processar_uma() is None
    pronta = f.obter(t1.id)
    assert set(pronta.arquivos()) == {"pdf", "planilha", "markdown"} and f.relatorio(t1.id).modo == "debater"
    assert [t.id for t in f.buscar("debentures")] == [t1.id]  # busca sem acento
    assert {t.id for t in f.a_entregar()} == {t1.id, t2.id}
    f.marcar_entregue(t1.id)
    assert [t.id for t in f.a_entregar()] == [t2.id]
    TIPOS.pop("teste_ok"), TIPOS.pop("teste_erro")


def _mercado() -> rf.Mercado:
    class Tit:  # título do Tesouro de mentira
        def __init__(self, nome, taxa, venc):
            self.nome, self.taxa_compra, self.vencimento = nome, taxa, venc

    return rf.Mercado(10.0, date(2026, 10, 2), {}, {2026: 4.0}, {"tesouro_prefixado": Tit("Tesouro Prefixado 2028", 12.0, date(2028, 1, 1))},
                      date(2026, 10, 2), ["📊 teste"], [])


def test_renda_fixa_conferida_a_mao():
    cen = rf.Cenario("plano", [10.0] * 12, [4.0] * 12)
    pre = rf.calcular(rf.Opcao("cdb_pre", "CDB pré 10%", taxa_pre=10.0), 100_000, 12, cen)
    assert pre.bruto == pytest.approx(110_000, abs=0.5)  # 100 mil × 1,10
    assert pre.aliquota == 0.175 and pre.liquido == pytest.approx(108_250, abs=0.5)  # 365 dias (361–720) → 17,5%
    assert pre.ganho_real == pytest.approx(108_250 - 104_000, abs=1)
    lci = rf.calcular(rf.Opcao("lci", "LCI", percentual_cdi=100), 100_000, 12, cen)
    assert lci.ir == 0 and lci.liquido == pytest.approx(100_000 * 1.099, abs=1)  # CDI = Selic − 0,10
    tes = rf.calcular(rf.Opcao("tesouro_prefixado", "TP", taxa_pre=10.0), 100_000, 12, cen)
    assert tes.custodia == pytest.approx(210, abs=5)  # 0,20% sobre o saldo médio do ano (~105 mil)
    # trajetória da Selic: 10% hoje → 8% no fim do ano seguinte (linear)
    c = rf.montar_cenario(15, 10.0, {2026: 9.0}, {2026: 4.0, 2027: 3.5}, hoje=date(2026, 9, 30))
    assert c.selic[0] > c.selic[2] > 9.0 > c.selic[5] or c.selic[3] == pytest.approx(9.0, abs=0.2)
    assert c.ipca[0] == 4.0 and c.ipca[-1] == 3.5


def test_comparativo_monta_relatorio_sem_internet():
    rel = rf.comparativo({"valor": 50_000, "prazo_anos": 1, "opcoes": [{"tipo": "cdb", "percentual_cdi": 110},
                          {"tipo": "lca", "percentual_cdi": 92}, {"tipo": "tesouro_prefixado"}, {"tipo": "xpto"}]},
                         mercado=_mercado(), redigir=False, hoje=date(2026, 10, 4))
    tab = rel.secoes[0].tabelas[0]
    assert [l[0] for l in tab.linhas][0] in {"CDB 110,0% do CDI", "Tesouro Prefixado 2028", "LCA 92,0% do CDI"}
    assert len(tab.linhas) == 3 and any("xpto" in a for a in rel.avisos)
    assert "equivale a um CDB de" in " ".join(rel.resumo) and rel.fatos["melhor"]
    assert redacao.conferir(" ".join(rel.resumo), redacao.referencias(rel)) == []


def test_calculadoras_novas_conferidas_a_mao():
    def linhas(nome, **kw):
        return dict(calc.executar(nome, kw).linhas)

    assert linhas("financiamento", valor=300000, taxa_am=1, meses=360)["Parcela fixa"] == "R$ 3.085,84"
    sac = linhas("financiamento", valor=100000, taxa_am=1, meses=12, sistema="sac")
    assert sac["1ª parcela"] == "R$ 9.333,33" and sac["Última parcela"] == "R$ 8.416,67" and sac["Total de juros"] == "R$ 6.500,00"
    v = linhas("vpl_tir", fluxos="-1000; 300; 400; 500", taxa_desconto=10)
    assert v["VPL"] == "R$ -21,04" and v["TIR"].startswith("8,9")
    assert linhas("pu_prefixado", taxa_aa=13.5, dias_uteis=504)["PU"] == "R$ 776,26"
    d = linhas("duration", fluxos="10; 10; 10; 110", taxa_aa=10)
    assert d["Duration de Macaulay"] == "3,49 anos" and d["Preço (VP dos fluxos)"] == "R$ 100,00"
    assert linhas("renda_aposentadoria", patrimonio=1_200_000, taxa_real_aa=0, anos=10)["Renda mensal"] == "R$ 10.000,00"
    assert linhas("aporte_necessario", meta=120_000, anos=10, taxa_aa=0)["Aporte mensal"] == "R$ 1.000,00"
    assert calc.executar("financiamento", {"valor": "100.000,00", "taxa_am": "1", "meses": "12"}).linhas[0][1] == "R$ 8.884,88"
    assert set(calc.ESQUEMAS) == set(calc.CALCULADORAS)


def test_terminal_relatorios_e_calculadoras(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    from quiron.servicos.analise import fila as mod
    from quiron.terminal.backend import app as terminal

    monkeypatch.setenv("QUIRON_DADOS", str(tmp_path))
    monkeypatch.delenv("TERMINAL_SENHA", raising=False)
    monkeypatch.setattr(mod, "_FILA", None)
    c = TestClient(terminal.app)
    lista = c.get("/api/calc").json()
    assert "financiamento" in lista and lista["financiamento"]["campos"][3][3] == "opção"
    assert c.post("/api/calc/vpl_tir", json={"fluxos": "-100; 110", "taxa_desconto": "0"}).json()["linhas"][0][1] == "R$ 10,00"
    assert c.get("/api/relatorios").json() == {"itens": []}
    assert c.get("/relatorios/1/relatorio.pdf").status_code == 404
    assert c.get("/relatorios/1/..%2F..%2Fsegredo").status_code == 404


def test_tarefa_orfa_volta_para_a_fila(tmp_path):
    f = Fila(tmp_path / "a.db", tmp_path / "rel")
    with f._con() as c:
        c.execute("INSERT INTO tarefas (tipo, parametros, situacao, criada_em, iniciada_em, dono) VALUES "
                  "('renda_fixa_comparativo', '{}', 'rodando', '2026-10-04T10:00:00', datetime('now','localtime'), '999999999')")
    assert f.recuperar_orfas() == 1 and f.obter(1).situacao == "na fila"
