"""Relatório do Quíron: um modelo único que vira PDF (capa, premissas, fontes, limitações), planilha e Markdown.

Usado por todas as análises (Fases 7 a 11). Os números chegam calculados em Python; o texto é redigido depois
(`redacao.py`) a partir deles. PDF com `pymupdf.Story` (HTML+CSS, sem dependência de sistema: funciona igual no
Windows e no Linux), gráficos com matplotlib e planilha com openpyxl.
"""

from __future__ import annotations

import html
import io
import json
import re
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

# Paleta validada (modo claro, impressão): ordem fixa das séries; série única em azul.
SERIES = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7", "#e34948"]
TEXTO, TEXTO_2, GRADE = "#0b0b0b", "#52514e", "#e4e3df"
RODAPE_PADRAO = "Uso interno e de estudo — não constitui recomendação nem relatório de análise (Resolução CVM 20)."
MODOS = {"entregar": "Entregar pronto", "debater": "Debate (prós, contras e perguntas)",
         "contestar": "Advogado do diabo (contestação)"}


# ---------------------------------------------------------------- formatação brasileira
def brl(v: float) -> str:
    return "R$ " + f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def pct(v: float, casas: int = 2) -> str:
    return f"{v:.{casas}f}".replace(".", ",") + "%"


def num(v: float, casas: int = 2) -> str:
    return f"{v:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def formatar(v: Any, formato: str) -> str:
    if v is None or v == "":
        return "—"
    if isinstance(v, str):
        return v
    return {"brl": brl, "pct": pct, "num": num, "int": lambda x: num(x, 0), "reais": lambda x: num(x, 0)}.get(formato, str)(v)


# ---------------------------------------------------------------- peças
@dataclass
class Tabela:
    titulo: str
    colunas: list[str]
    linhas: list[list[Any]]
    formatos: list[str] = field(default_factory=list)  # por coluna: texto | brl | pct | num | int
    nota: str = ""

    def formato(self, i: int) -> str:
        return self.formatos[i] if i < len(self.formatos) else "texto"


@dataclass
class Grafico:
    titulo: str
    tipo: str  # "barras_h" (categorias; 1 série ou agrupadas) | "linhas" (evolução) | "dispersao" (x × y)
    rotulos: list[str] = field(default_factory=list)  # barras: categorias · linhas: eixo x · dispersão: não usado
    series: dict[str, list[float]] = field(default_factory=dict)  # barras/linhas: {nome: valores} · dispersão: y
    formato: str = "num"  # dos valores (rótulos diretos e eixo y)
    eixo_x: str = ""
    nota: str = ""
    x: dict[str, list[float]] = field(default_factory=dict)  # dispersão: valores de x por série
    formato_x: str = "num"  # dispersão: formato do eixo x


@dataclass
class Secao:
    titulo: str
    texto: str = ""  # Markdown simples
    tabelas: list[Tabela] = field(default_factory=list)
    graficos: list[Grafico] = field(default_factory=list)


@dataclass
class Relatorio:
    titulo: str
    tipo: str
    subtitulo: str = ""
    modo: str = "entregar"
    resumo: list[str] = field(default_factory=list)
    premissas: list[str] = field(default_factory=list)
    secoes: list[Secao] = field(default_factory=list)
    fontes: list[str] = field(default_factory=list)
    limitacoes: list[str] = field(default_factory=list)
    fatos: dict[str, Any] = field(default_factory=dict)  # números calculados (base da redação e da conferência)
    parametros: dict[str, Any] = field(default_factory=dict)
    rodape: str = RODAPE_PADRAO
    criado_em: str = field(default_factory=lambda: datetime.now().isoformat(timespec="seconds"))
    avisos: list[str] = field(default_factory=list)

    # ------------------------------------------------------------ Markdown
    def markdown(self) -> str:
        p = [f"# {self.titulo}"]
        if self.subtitulo:
            p.append(f"_{self.subtitulo}_")
        p.append(f"{_data(self.criado_em)} · modo: {MODOS.get(self.modo, self.modo)}")
        if self.resumo:
            p += ["## Resumo", *[f"- {r}" for r in self.resumo]]
        if self.premissas:
            p += ["## Premissas", *[f"- {r}" for r in self.premissas]]
        for s in self.secoes:
            p.append(f"## {s.titulo}")
            if s.texto:
                p.append(s.texto)
            for t in s.tabelas:
                p.append(f"**{t.titulo}**\n")
                p.append("| " + " | ".join(t.colunas) + " |")
                p.append("|" + "---|" * len(t.colunas))
                p += ["| " + " | ".join(formatar(v, t.formato(i)) for i, v in enumerate(l)) + " |" for l in t.linhas]
                if t.nota:
                    p.append(f"_{t.nota}_")
        if self.fontes:
            p += ["## Fontes", *[f"- {f}" for f in self.fontes]]
        if self.limitacoes:
            p += ["## Limitações", *[f"- {l}" for l in self.limitacoes]]
        p.append(f"\n_{self.rodape}_")
        return "\n\n".join(p)

    def resumo_curto(self, limite: int = 1500) -> str:
        """Texto para o Telegram (o resto vai no PDF)."""
        linhas = [f"📑 {self.titulo}", *[f"• {r}" for r in self.resumo]]
        if self.avisos:
            linhas += [f"⚠️ {a}" for a in self.avisos]
        linhas.append(f"_{self.rodape}_")
        texto = "\n".join(linhas)
        return texto if len(texto) <= limite else texto[:limite - 1] + "…"

    # ------------------------------------------------------------ gráficos
    def graficos(self) -> list[tuple[str, Grafico]]:
        return [(f"g{i}_{j}.png", g) for i, s in enumerate(self.secoes) for j, g in enumerate(s.graficos)]

    # ------------------------------------------------------------ PDF
    def pdf(self, destino: Path) -> Path:
        import pymupdf

        arquivo = pymupdf.Archive()
        imagens = {}
        for nome, g in self.graficos():
            png = desenhar(g)
            arquivo.add(png, nome)
            pix = pymupdf.Pixmap(png)
            largura = 487
            imagens[id(g)] = (nome, largura, round(largura * pix.height / pix.width))
        escritor = pymupdf.DocumentWriter(str(destino))
        pagina = pymupdf.paper_rect("a4")
        esq, topo, dir_, base = 54, 56, pagina.width - 54, pagina.height - 60
        estado = {"disp": None, "y": topo}

        def nova_pagina() -> None:
            if estado["disp"] is not None:
                escritor.end_page()
            estado["disp"], estado["y"] = escritor.begin_page(pagina), topo

        nova_pagina()
        # Cada bloco é colocado à parte: gráficos (altura conhecida) nunca são encolhidos — se não cabem, vão para a
        # próxima página; texto e tabelas fluem normalmente entre páginas.
        for trecho, altura_min in self._blocos(imagens):
            story = pymupdf.Story(trecho, user_css=CSS, archive=arquivo)
            while True:
                if base - estado["y"] < 24 or (altura_min and base - estado["y"] < altura_min and estado["y"] > topo):
                    nova_pagina()
                    continue
                mais, preenchido = story.place(pymupdf.Rect(esq, estado["y"], dir_, base))
                story.draw(estado["disp"])
                estado["y"] = pymupdf.Rect(preenchido).y1 + 4
                if not mais:
                    break
                nova_pagina()
        escritor.end_page()
        escritor.close()
        _carimbar(destino, self.titulo, self.rodape)
        return destino

    def _blocos(self, imagens: dict[int, tuple[str, int, int]]) -> list[tuple[str, int]]:
        """(html, altura mínima): se sobrar menos que isso na página, o bloco começa na próxima (evita título órfão
        e gráfico encolhido). Texto e tabelas continuam fluindo entre páginas."""
        e = html.escape
        blocos: list[tuple[str, int]] = []
        capa = [f'<div class="capa"><p class="marca">QUÍRON · ANÁLISE</p><h1>{e(self.titulo)}</h1>']
        if self.subtitulo:
            capa.append(f'<p class="sub">{e(self.subtitulo)}</p>')
        capa.append(f'<p class="meta">{_data(self.criado_em)} · modo: {e(MODOS.get(self.modo, self.modo))}</p></div>')
        if self.avisos:
            capa.append('<div class="aviso">' + "<br/>".join(f"⚠ {e(a)}" for a in self.avisos) + "</div>")
        blocos.append(("".join(capa), 0))

        def lista(titulo: str, itens: list[str]) -> None:
            if itens:
                blocos.append((f"<h2>{e(titulo)}</h2><ul>" + "".join(f"<li>{_md_inline(i)}</li>" for i in itens) + "</ul>", 70))

        lista("Resumo", self.resumo)
        lista("Premissas", self.premissas)
        for s in self.secoes:
            cabeca = f"<h2>{e(s.titulo)}</h2>"  # o título da seção vai junto do primeiro bloco dela
            if s.texto:
                blocos.append((cabeca + md_para_html(s.texto), 70))
                cabeca = ""
            for t in s.tabelas:
                blocos.append((cabeca + _tabela_html(t), 90))
                cabeca = ""
            for g in s.graficos:
                nome, larg, alt = imagens[id(g)]
                nota = f'<p class="nota">{e(g.nota)}</p>' if g.nota else ""
                blocos.append((cabeca + f'<p class="titulo-graf">{e(g.titulo)}</p><img src="{nome}" width="{larg}" '
                               f'height="{alt}"/>{nota}', alt + (60 if cabeca else 34)))
                cabeca = ""
            if cabeca:
                blocos.append((cabeca, 0))
        lista("Fontes", self.fontes)
        lista("Limitações", self.limitacoes)
        return blocos

    # ------------------------------------------------------------ planilha
    def planilha(self, destino: Path) -> Path:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Resumo"
        negrito = Font(bold=True)
        ws.append([self.titulo])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([f"{_data(self.criado_em)} · modo: {MODOS.get(self.modo, self.modo)}"])
        for titulo, itens in (("Resumo", self.resumo), ("Premissas", self.premissas), ("Fontes", self.fontes),
                              ("Limitações", self.limitacoes)):
            if itens:
                ws.append([])
                ws.append([titulo])
                ws.cell(ws.max_row, 1).font = negrito
                for i in itens:
                    ws.append([_sem_md(i)])
        ws.append([])
        ws.append([self.rodape])
        ws.column_dimensions["A"].width = 120
        for linha in ws.iter_rows():
            for c in linha:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        formatos = {"brl": '"R$" #,##0.00', "reais": '"R$" #,##0.00', "pct": '0.00"%"', "num": "#,##0.00", "int": "#,##0"}
        usados: set[str] = {"Resumo"}
        cabecalho = PatternFill("solid", fgColor="EEEEEC")
        for s in self.secoes:
            for t in s.tabelas:
                nome = _nome_aba(t.titulo, usados)
                aba = wb.create_sheet(nome)
                aba.append(t.colunas)
                for c in aba[1]:
                    c.font, c.fill = negrito, cabecalho
                for l in t.linhas:
                    aba.append(l)
                for i in range(len(t.colunas)):
                    fmt = formatos.get(t.formato(i))
                    if fmt:
                        for (c,) in aba.iter_rows(min_row=2, min_col=i + 1, max_col=i + 1):
                            if isinstance(c.value, (int, float)):
                                c.number_format = fmt
                    aba.column_dimensions[get_column_letter(i + 1)].width = max(14, min(48, len(str(t.colunas[i])) + 6))
                if t.nota:
                    aba.append([])
                    aba.append([t.nota])
            for g in s.graficos:
                nome = _nome_aba("Dados - " + g.titulo, usados)
                aba = wb.create_sheet(nome)
                if g.tipo == "linhas":
                    aba.append([g.eixo_x or "x", *g.series.keys()])
                    for k, rot in enumerate(g.rotulos):
                        aba.append([rot, *[v[k] if k < len(v) else None for v in g.series.values()]])
                elif g.tipo == "dispersao":
                    aba.append(["Série", g.eixo_x or "x", "y"])
                    for nome, ys in g.series.items():
                        for xv, yv in zip(g.x.get(nome, []), ys):
                            aba.append([nome, xv, yv])
                else:
                    aba.append(["Categoria", *g.series.keys()])
                    for k, rot in enumerate(g.rotulos):
                        aba.append([rot, *[v[k] if k < len(v) else None for v in g.series.values()]])
                for c in aba[1]:
                    c.font, c.fill = negrito, cabecalho
        wb.save(destino)
        return destino

    # ------------------------------------------------------------ dados
    def como_dict(self) -> dict[str, Any]:
        return asdict(self)

    @classmethod
    def de_dict(cls, d: dict[str, Any]) -> "Relatorio":
        d = dict(d)
        d["secoes"] = [Secao(s["titulo"], s.get("texto", ""), [Tabela(**t) for t in s.get("tabelas", [])],
                             [Grafico(**g) for g in s.get("graficos", [])]) for s in d.get("secoes", [])]
        return cls(**d)

    def salvar(self, pasta: Path) -> dict[str, Path]:
        """Grava PDF, planilha, Markdown e o JSON (para reabrir) na pasta. Devolve os caminhos."""
        pasta.mkdir(parents=True, exist_ok=True)
        caminhos = {"pdf": self.pdf(pasta / "relatorio.pdf"), "planilha": self.planilha(pasta / "planilha.xlsx"),
                    "markdown": pasta / "relatorio.md", "json": pasta / "relatorio.json"}
        caminhos["markdown"].write_text(self.markdown(), encoding="utf-8")
        caminhos["json"].write_text(json.dumps(self.como_dict(), ensure_ascii=False, indent=1, default=str), encoding="utf-8")
        return caminhos


# ---------------------------------------------------------------- auxiliares
CSS = f"""
body {{ font-family: sans-serif; font-size: 10pt; color: {TEXTO}; line-height: 1.35; }}
.capa {{ border-bottom: 2px solid #f0a020; padding-bottom: 8pt; margin-bottom: 10pt; }}
.marca {{ color: #a26a0b; font-size: 8pt; letter-spacing: 1pt; font-weight: bold; margin: 0; }}
h1 {{ font-size: 20pt; margin: 4pt 0 2pt 0; }}
.sub {{ font-size: 11pt; color: {TEXTO_2}; margin: 0; }}
.meta {{ font-size: 8.5pt; color: {TEXTO_2}; margin: 4pt 0 0 0; }}
h2 {{ font-size: 12.5pt; margin: 14pt 0 4pt 0; color: {TEXTO}; border-bottom: 0.5pt solid {GRADE}; padding-bottom: 2pt; }}
ul {{ margin: 2pt 0 2pt 0; }}
li {{ margin-bottom: 2pt; }}
table {{ border-collapse: collapse; width: 100%; margin: 6pt 0 2pt 0; }}
th {{ font-size: 8.5pt; text-align: left; padding: 3pt 5pt; border-bottom: 0.9pt solid #8d8c84; }}
td {{ font-size: 8.5pt; padding: 3pt 5pt; border-bottom: 0.4pt solid {GRADE}; }}
table.densa th, table.densa td {{ font-size: 7.5pt; padding: 2.5pt 3pt; }}
td.n {{ text-align: right; }}
.titulo-tab, .titulo-graf {{ font-weight: bold; font-size: 9.5pt; margin: 10pt 0 2pt 0; }}
.nota {{ font-size: 7.5pt; color: {TEXTO_2}; margin: 2pt 0 6pt 0; }}
.aviso {{ border-left: 3pt solid #eda100; padding: 2pt 6pt; font-size: 8.5pt; margin: 6pt 0; color: #5a4300; }}
"""


def _data(iso: str) -> str:
    try:
        return datetime.fromisoformat(iso).strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return iso


EMOJI = re.compile("[\U0001F300-\U0001FAFF\u2600-\u27BF\uFE0F]")


def _md_inline(texto: str) -> str:
    t = html.escape(EMOJI.sub("", texto).strip())
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"(?<![\w*])_(.+?)_(?![\w*])", r"<i>\1</i>", t)
    return t


def _sem_md(texto: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", texto)


def md_para_html(texto: str) -> str:
    """Markdown simples (parágrafos, listas, negrito/itálico, ### títulos) → HTML para o PDF."""
    import markdown

    return markdown.markdown(texto, extensions=["tables"], output_format="html")


def _tabela_html(t: Tabela) -> str:
    e = html.escape
    numericos = [t.formato(i) in {"brl", "pct", "num", "int", "reais"} for i in range(len(t.colunas))]
    cab = "".join(f'<th style="text-align:{"right" if n else "left"}">{e(c)}</th>' for c, n in zip(t.colunas, numericos))
    corpo = "".join("<tr>" + "".join(f'<td class="{"n" if numericos[i] else ""}">{e(formatar(v, t.formato(i)))}</td>'
                                     for i, v in enumerate(l)) + "</tr>" for l in t.linhas)
    nota = f'<p class="nota">{e(t.nota)}</p>' if t.nota else ""
    classe = ' class="densa"' if len(t.colunas) >= 7 else ""
    return f'<p class="titulo-tab">{e(t.titulo)}</p><table{classe}><tr>{cab}</tr>{corpo}</table>{nota}'


def _nome_aba(titulo: str, usados: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "", titulo)[:28] or "Tabela"
    nome, i = base, 2
    while nome in usados:
        nome = f"{base[:25]} {i}"
        i += 1
    usados.add(nome)
    return nome


def _latin1(texto: str) -> str:
    """As fontes-padrão do PDF (cabeçalho/rodapé) só têm Latin-1: troca travessões e aspas tipográficas."""
    trocas = {"—": "-", "–": "-", "−": "-", "“": '"', "”": '"', "‘": "'", "’": "'", "…": "...", "≈": "~"}
    texto = "".join(trocas.get(c, c) for c in texto)
    return texto.encode("latin-1", "ignore").decode("latin-1")


def _carimbar(caminho: Path, titulo: str, rodape: str) -> None:
    """Cabeçalho discreto + rodapé de compliance e numeração em todas as páginas."""
    import pymupdf

    titulo, rodape = _latin1(titulo), _latin1(rodape)

    doc = pymupdf.open(caminho)
    total = doc.page_count
    for i, pg in enumerate(doc, 1):
        r = pg.rect
        if i > 1:
            pg.insert_textbox(pymupdf.Rect(54, 24, r.width - 54, 40), titulo[:110], fontsize=7, color=(0.45, 0.45, 0.43))
        pg.insert_textbox(pymupdf.Rect(54, r.height - 44, r.width - 110, r.height - 20), rodape, fontsize=6.5,
                          color=(0.45, 0.45, 0.43))
        pg.insert_textbox(pymupdf.Rect(r.width - 110, r.height - 44, r.width - 54, r.height - 20), f"{i}/{total}",
                          fontsize=7, color=(0.45, 0.45, 0.43), align=2)
    # o Story grava as imagens sem compressão: salvar de novo comprimido (gráficos de ~1 MB cada viram dezenas de KB)
    temporario = caminho.with_suffix(".tmp.pdf")
    doc.save(temporario, garbage=4, deflate=True, deflate_images=True, deflate_fonts=True)
    doc.close()
    temporario.replace(caminho)


def desenhar(g: Grafico) -> bytes:
    """Gráfico em PNG (impressão, modo claro): marcas finas, grade discreta, rótulos diretos, legenda se ≥2 séries."""
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams.update({"font.size": 9, "axes.edgecolor": GRADE, "axes.labelcolor": TEXTO_2, "xtick.color": TEXTO_2,
                         "ytick.color": TEXTO_2, "font.family": "DejaVu Sans"})
    fmt = (lambda v: formatar(v, g.formato))
    if g.tipo == "barras_h":
        _barras(plt, g, fmt)
        fig, ax = plt.gcf(), plt.gca()
    elif g.tipo == "dispersao":
        fig, ax = plt.subplots(figsize=(6.4, 3.8), dpi=200)
        k_ponto = 0
        for k, (nome, ys) in enumerate(g.series.items()):
            xs = g.x.get(nome, [])
            cor = SERIES[k % len(SERIES)]
            if len(ys) > 2:  # curva (ex.: fronteira eficiente)
                ax.plot(xs, ys, color=cor, linewidth=1.6, label=nome, zorder=2)
            else:  # pontos destacados com rótulo direto
                ax.scatter(xs, ys, s=46, color=cor, edgecolor="white", linewidth=1.5, label=nome, zorder=3)
                for xv, yv in zip(xs, ys):
                    ax.annotate(nome, (xv, yv), xytext=(6, 6 if k_ponto % 2 == 0 else -10), textcoords="offset points",
                                fontsize=7.5, color=TEXTO)
                k_ponto += 1
        ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: formatar(v, g.formato_x)))
        ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: _eixo(v, g.formato)))
        ax.grid(color=GRADE, linewidth=0.6)
        ax.set_axisbelow(True)
        for lado in ("top", "right"):
            ax.spines[lado].set_visible(False)
        if g.eixo_x:
            ax.set_xlabel(g.eixo_x)
        if len(g.series) >= 2:
            ax.legend(frameon=False, fontsize=8, loc="upper center", bbox_to_anchor=(0.5, -0.16),
                      ncol=min(4, len(g.series)))
        ax.margins(x=0.12, y=0.12)
    else:
        fig, ax = plt.subplots(figsize=(6.4, 3.6), dpi=200)
        x = list(range(len(g.rotulos)))
        rotular = len(g.series) <= 4  # rótulo direto no fim da linha só até 4 séries (senão, só a legenda)
        for k, (nome, valores) in enumerate(g.series.items()):
            cor = SERIES[k % len(SERIES)]
            ax.plot(x[:len(valores)], valores, color=cor, linewidth=1.6, label=nome)
            if valores and rotular:
                ax.annotate(fmt(valores[-1]), (x[len(valores) - 1], valores[-1]), xytext=(4, 0),
                            textcoords="offset points", fontsize=8, color=TEXTO, va="center")
        passo = max(1, len(x) // 8)
        ax.set_xticks(x[::passo], [g.rotulos[i] for i in x[::passo]])
        ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: _eixo(v, g.formato)))
        ax.grid(axis="y", color=GRADE, linewidth=0.6)
        ax.set_axisbelow(True)
        for lado in ("top", "right"):
            ax.spines[lado].set_visible(False)
        if g.eixo_x:
            ax.set_xlabel(g.eixo_x)
        if len(g.series) >= 2:
            ax.legend(frameon=False, fontsize=8, loc="upper center", bbox_to_anchor=(0.5, -0.16),
                      ncol=min(3, len(g.series)))
        ax.margins(x=0.02)
        fig.subplots_adjust(right=0.78)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor="white", bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)
    return buf.getvalue()


def _barras(plt, g: Grafico, fmt) -> None:
    """Barras horizontais: série única (azul) ou agrupadas (ordem fixa da paleta + legenda). Aceita negativos."""
    nomes = list(g.series)
    n_cat, n_ser = len(g.rotulos), max(1, len(nomes))
    todos = [v for vs in g.series.values() for v in vs if v is not None] or [0]
    menor, maior = min(0, min(todos)), max(0, max(todos))
    amplitude = (maior - menor) or 1
    altura_barra = 0.55 if n_ser == 1 else min(0.8 / n_ser, 0.3)
    altura = max(1.6, (0.42 if n_ser == 1 else 0.22 * n_ser + 0.2) * n_cat + 0.6)
    fig, ax = plt.subplots(figsize=(6.4, altura), dpi=200)
    base = list(range(n_cat))[::-1]
    for k, nome in enumerate(nomes):
        desloc = (k - (n_ser - 1) / 2) * altura_barra
        pos = [b - desloc for b in base]
        valores = [v if v is not None else 0 for v in g.series[nome]]
        ax.barh(pos, valores, color=SERIES[k % len(SERIES)], height=altura_barra * 0.92, edgecolor="white",
                linewidth=1, label=nome)
        if n_ser <= 4:
            for y, v in zip(pos, valores):
                lado = 1 if v >= 0 else -1
                ax.text(v + lado * amplitude * 0.01, y, fmt(v), va="center", ha="left" if v >= 0 else "right",
                        fontsize=7 if n_ser > 1 else 7.5, color=TEXTO)
    ax.set_yticks(base, g.rotulos)
    ax.set_xlim(menor - (amplitude * 0.22 if menor < 0 else 0), maior + (amplitude * 0.22 if maior > 0 else amplitude * 0.02))
    if menor < 0:
        ax.axvline(0, color="#8d8c84", linewidth=0.8)
    ax.xaxis.set_visible(False)
    for lado in ("top", "right", "bottom"):
        ax.spines[lado].set_visible(False)
    if n_ser >= 2:
        ax.legend(frameon=False, fontsize=8, loc="upper center", bbox_to_anchor=(0.5, 0), ncol=min(4, n_ser))


def _eixo(v: float, formato: str) -> str:
    """Rótulo do eixo: valores em reais ficam compactos (R$ 850 mil, R$ 6,2 mi); o resto segue o formato."""
    if formato in {"brl", "reais"} and abs(v) >= 10_000:
        if abs(v) >= 1_000_000:
            return "R$ " + f"{v / 1_000_000:.1f}".replace(".", ",").replace(",0", "") + " mi"
        return f"R$ {v / 1000:.0f} mil"
    return formatar(v, formato)
