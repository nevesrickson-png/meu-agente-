"""Leitura da carteira a partir de texto livre, planilha (xlsx/csv) ou print (OCR local com tesseract).

- Texto: o cérebro estrutura em JSON (cliente só como CLI-XXX); se falhar, um leitor por regras cobre linhas simples.
- Planilha: reconhece colunas comuns (ativo, valor, quantidade, ticker, classe, custo, vencimento, taxa…).
- Print: OCR no próprio computador; CPF, CNPJ, conta, agência, e-mail e telefone são mascarados ANTES de ir ao modelo.
- Posição com quantidade e ticker ganha o valor pela cotação do dia.
"""

from __future__ import annotations

import csv
import io
import json
import logging
import re
import shutil
import subprocess
import tempfile
from datetime import date
from pathlib import Path
from typing import Any

from quiron.nucleo import cerebro
from quiron.servicos.carteira.modelo import CLASSES, Carteira, Posicao, _data, classe_por_ticker

RE_TICKER = re.compile(r"\b([A-Z]{4}\d{1,2})\b")
RE_VALOR = re.compile(r"R\$\s*([\d.]+(?:,\d{1,2})?)\s*(mil|mi|milh[õo]es)?|([\d.]+(?:,\d{1,2})?)\s*(mil|mi)\b", re.I)


def classe_por_nome(nome: str, ticker: str = "") -> str:
    """Classe provável pelo nome do produto (palavras-chave do mercado brasileiro)."""
    n = nome.lower()
    if ticker and (c := classe_por_ticker(ticker)):
        return c
    regras = [("internacional", ("ivvb", "bdr", "exterior", "global", "s&p", "nasdaq", "offshore", "internacional")),
              ("cripto", ("bitcoin", "btc", "cripto", "ethereum", "hash11")),
              ("inflacao", ("ipca", "ntn-b", "ntnb", "inflação", "inflacao", "ima-b", "imab", "renda+", "educa+")),
              ("prefixado", ("prefixad", "pré-fixad", "pre-fixad", "ltn", "ntn-f", "irf-m", "% a.a. pré", "pré ")),
              ("fii", ("fii", "fundo imobili", "imobiliário")),
              ("multimercado", ("multimercado", "macro", "long short", "long-short", "multi ")),
              ("acoes", ("ações", "acoes", "ação", "fundo de ações", "fia", "small caps", "dividendos", "ibovespa", "bova")),
              ("pos_fixado", ("cdi", "selic", "di ", "pós", "pos-fix", "cdb", "lci", "lca", "lc ", "poupança", "caixa",
                              "conta remunerada", "tesouro selic", "lft", "debênture", "debenture", "cri", "cra"))]
    for classe, chaves in regras:
        if any(k in n + " " for k in chaves):
            return classe
    return "pos_fixado"


def tipo_por_nome(nome: str) -> str:
    n = nome.lower()
    for t, chaves in (("lci", ("lci",)), ("lca", ("lca",)), ("cri", ("cri ",)), ("cra", ("cra ",)),
                      ("debenture_incentivada", ("incentivad",)), ("debenture", ("debênture", "debenture")),
                      ("cdb", ("cdb",)), ("tesouro_selic", ("tesouro selic", "lft")), ("tesouro_ipca", ("tesouro ipca", "ntn-b")),
                      ("tesouro_prefixado", ("tesouro prefixado", "ltn")), ("previdencia", ("pgbl", "vgbl", "previdência")),
                      ("fundo", ("fundo", "fic", "fia")), ("poupanca", ("poupança",))):
        if any(k in n + " " for k in chaves):
            return t
    return ""


def numero(texto: Any) -> float | None:
    """'R$ 1.234,56' · '50 mil' · '1,2 mi' · 1234.5 → float."""
    if texto is None or texto == "":
        return None
    if isinstance(texto, (int, float)):
        return float(texto)
    t = str(texto).lower().replace("r$", "").strip()
    mult = 1.0
    if re.search(r"\bmi\b|\bmilh", t):
        mult = 1e6
    elif "mil" in t:
        mult = 1e3
    t = re.sub(r"[^\d,.\-]", "", t)
    if not t:
        return None
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    elif t.count(".") > 1 or re.fullmatch(r"-?\d{1,3}(\.\d{3})+", t):
        t = t.replace(".", "")
    try:
        return float(t) * mult
    except ValueError:
        return None


def mascarar_identificadores(texto: str) -> tuple[str, int]:
    """Esconde CPF, CNPJ, conta/agência, e-mail e telefone (para nunca mandar dado identificável ao modelo)."""
    padroes = [r"\b\d{3}\.?\d{3}\.?\d{3}-?\d{2}\b", r"\b\d{2}\.?\d{3}\.?\d{3}/?\d{4}-?\d{2}\b",
               r"(?i)\b(conta|ag[êe]ncia|ag\.|c/c|cc)\s*[:nº°#]*\s*[\d\-.]{3,}", r"[\w.+-]+@[\w-]+\.[\w.]+",
               r"\(?\b\d{2}\)?\s?9?\d{4}-?\d{4}\b"]
    total = 0
    for p in padroes:
        texto, n = re.subn(p, "[oculto]", texto)
        total += n
    return texto, total


# ---------------------------------------------------------------- texto
SISTEMA = ("Você estrutura carteiras de investimento brasileiras em JSON. Não invente posições nem valores: só o que está "
           "no texto. Cliente apenas como código CLI-XXX (nunca nome).")


def _pedido(texto: str) -> str:
    classes = ", ".join(f"{k} ({v['nome']})" for k, v in CLASSES.items())
    return (f"Texto da carteira:\n<<<\n{texto[:8000]}\n>>>\n\nClasses válidas: {classes}.\n"
            "Responda SOMENTE com JSON: {\"nome\": \"...\", \"cliente\": \"CLI-XXX ou vazio\", \"perfil\": "
            "\"conservador|moderado|arrojado|vazio\", \"posicoes\": [{\"nome\": \"...\", \"classe\": \"...\", "
            "\"valor\": número em reais ou null, \"ticker\": \"PETR4 ou vazio\", \"quantidade\": número ou null, "
            "\"tipo\": \"cdb|lci|lca|cri|cra|debenture|debenture_incentivada|tesouro_selic|tesouro_ipca|tesouro_prefixado|"
            "acao|fii|etf|fundo|previdencia|poupanca|outro\", \"custo\": número ou null, \"data_aplicacao\": \"AAAA-MM-DD ou vazio\", "
            "\"vencimento\": \"AAAA-MM-DD ou AAAA ou vazio\", \"taxa\": \"texto como 110% CDI, IPCA+6,5%, 13% a.a.\"}]}\n"
            "Valores como '50 mil' = 50000; '1,2 mi' = 1200000. Fundo de ações → acoes; previdência → classe do que ela "
            "investe (se não souber, multimercado).")


def _do_json(dados: dict[str, Any]) -> Carteira:
    posicoes, avisos = [], []
    for p in dados.get("posicoes") or []:
        nome = str(p.get("nome") or p.get("ticker") or "").strip()
        if not nome:
            continue
        ticker = str(p.get("ticker") or "").upper().strip()
        if not ticker and (m := RE_TICKER.search(nome.upper())) and classe_por_ticker(m.group(1)):
            ticker = m.group(1)  # "HGLG11" ou "PETR4 - Petrobras" no nome: vira o ticker
        classe = str(p.get("classe") or "").strip() or classe_por_nome(nome, ticker)
        if classe not in CLASSES:
            classe = classe_por_nome(nome, ticker)
        try:
            posicoes.append(Posicao(
                nome=nome, classe=classe, valor=numero(p.get("valor")) or 0.0, ticker=ticker,
                quantidade=numero(p.get("quantidade")), tipo=str(p.get("tipo") or tipo_por_nome(nome)),
                custo=numero(p.get("custo")), data_aplicacao=_data(str(p.get("data_aplicacao") or "")),
                vencimento=_data(str(p.get("vencimento") or "")), taxa=str(p.get("taxa") or "")))
        except (ValueError, TypeError) as e:
            avisos.append(f"Posição ignorada ({nome}): {e}")
    cliente = str(dados.get("cliente") or "")
    if cliente and not re.fullmatch(r"CLI-\w+", cliente):
        avisos.append("Cliente identificado só por código (CLI-XXX); o nome não foi guardado.")
        cliente = ""
    perfil = str(dados.get("perfil") or "").lower().strip()
    perfil = perfil if perfil in {"conservador", "moderado", "arrojado"} else ""
    return Carteira(str(dados.get("nome") or "Carteira"), posicoes, perfil, cliente, avisos)


def ler_por_regras(texto: str) -> Carteira:
    """Leitor simples, sem IA: uma posição por linha ('PETR4 200', 'CDB 110% CDI R$ 50 mil venc 2027')."""
    posicoes = []
    for linha in re.split(r"[\n;]+", texto):
        linha = linha.strip(" -•*\t")
        if len(linha) < 3:
            continue
        ticker = (RE_TICKER.search(linha.upper()) or [None, ""])[1] if RE_TICKER.search(linha.upper()) else ""
        valor = None
        if m := RE_VALOR.search(linha):
            valor = numero(m.group(0))
        quantidade = None
        if ticker and valor is None and (m := re.search(rf"{ticker}\s*[:\-]?\s*(\d+(?:[.,]\d+)?)", linha.upper())):
            quantidade = numero(m.group(1))
        venc = None
        if m := re.search(r"(?:venc\w*|até|ate)\s*(\d{1,2}/\d{4}|\d{2}/\d{2}/\d{4}|\d{4})", linha, re.I):
            venc = _data(m.group(1))
        if valor is None and quantidade is None:
            continue
        nome = re.sub(r"\s+", " ", linha)[:80]
        posicoes.append(Posicao(nome=nome, classe=classe_por_nome(nome, ticker), valor=valor or 0.0, ticker=ticker,
                                quantidade=quantidade, tipo=tipo_por_nome(nome), vencimento=venc,
                                taxa=(re.search(r"\d+[,.]?\d*\s*%[^;,\n]*", linha) or [""])[0] if "%" in linha else ""))
    return Carteira("Carteira", posicoes)


def ler_texto(texto: str, usar_cerebro: bool = True, **extras: Any) -> Carteira:
    texto, mascarados = mascarar_identificadores(texto)
    carteira = None
    if usar_cerebro:
        try:
            r = cerebro.perguntar(_pedido(texto), sistema=SISTEMA, temperatura=0, max_tokens=8000,
                                  response_format={"type": "json_object"}, **extras)
            bruto = re.sub(r"^```(?:json)?\s*|\s*```$", "", r.texto.strip())
            carteira = _do_json(json.loads(bruto[bruto.find("{"):bruto.rfind("}") + 1]))
        except (cerebro.CerebroIndisponivel, ValueError, json.JSONDecodeError) as e:
            logging.warning("leitura da carteira pelo cérebro falhou (%s); usando regras", type(e).__name__)
    if carteira is None or not carteira.posicoes:
        carteira = ler_por_regras(texto)
        carteira.avisos.append("Leitura por regras simples (sem IA): confira classes e valores.")
    if mascarados:
        carteira.avisos.append(f"{mascarados} dado(s) pessoal(is) (CPF/conta/e-mail/telefone) foram ocultados antes da leitura.")
    return carteira


# ---------------------------------------------------------------- planilha
COLUNAS = {"nome": ("ativo", "nome", "produto", "papel", "descrição", "descricao", "investimento", "título", "titulo"),
           "valor": ("valor", "saldo", "posição", "posicao", "valor atual", "saldo bruto", "total", "valor bruto"),
           "quantidade": ("quantidade", "qtd", "qtde", "cotas"), "ticker": ("ticker", "código", "codigo", "symbol"),
           "classe": ("classe", "categoria", "tipo de ativo", "estratégia", "estrategia"),
           "custo": ("custo", "valor aplicado", "aplicado", "investido", "valor investido", "custo total"),
           "data_aplicacao": ("data aplicação", "data aplicacao", "data de aplicação", "aplicação", "data compra"),
           "vencimento": ("vencimento", "venc", "data vencimento"), "taxa": ("taxa", "rentabilidade contratada", "indexador")}


def _mapear(cabecalho: list[str]) -> dict[str, int]:
    mapa = {}
    norm = [str(c or "").strip().lower() for c in cabecalho]
    for campo, nomes in COLUNAS.items():
        for i, c in enumerate(norm):
            if c in nomes or any(c.startswith(n) for n in nomes):
                mapa.setdefault(campo, i)
                break
    return mapa


def ler_planilha(conteudo: bytes, nome_arquivo: str) -> Carteira:
    if nome_arquivo.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        linhas = [list(r) for r in wb.active.iter_rows(values_only=True)]
    else:
        texto = conteudo.decode("utf-8-sig", errors="replace")
        sep = ";" if texto.count(";") >= texto.count(",") else ","
        linhas = list(csv.reader(io.StringIO(texto), delimiter=sep))
    linhas = [l for l in linhas if any(str(c or "").strip() for c in l)]
    cab_i = next((i for i, l in enumerate(linhas[:15]) if {"nome", "valor"} <= set(_mapear(l)) or
                  {"nome", "quantidade"} <= set(_mapear(l))), None)
    if cab_i is None:
        raise ValueError("Não achei o cabeçalho (preciso de uma coluna de ativo/nome e uma de valor ou quantidade).")
    mapa = _mapear(linhas[cab_i])
    dados = {"nome": Path(nome_arquivo).stem, "posicoes": []}
    for l in linhas[cab_i + 1:]:
        item = {campo: (l[i] if i < len(l) else None) for campo, i in mapa.items()}
        if not str(item.get("nome") or "").strip() or str(item.get("nome")).strip().lower() in {"total", "totais"}:
            continue
        for k in ("data_aplicacao", "vencimento"):
            v = item.get(k)
            if isinstance(v, date):
                item[k] = v.isoformat()[:10]
        dados["posicoes"].append(item)
    return _do_json(dados)


# ---------------------------------------------------------------- print (imagem)
def ocr(imagem: bytes) -> str:
    """OCR local (tesseract, português+inglês). Nada sai do computador nesta etapa."""
    if not shutil.which("tesseract"):
        raise RuntimeError("OCR indisponível neste computador (tesseract não instalado): mande como texto ou planilha.")
    with tempfile.TemporaryDirectory() as pasta:
        arq = Path(pasta) / "print.png"
        arq.write_bytes(imagem)
        r = subprocess.run(["tesseract", str(arq), "stdout", "-l", "por+eng", "--psm", "6"], capture_output=True,
                           text=True, timeout=120)
        if r.returncode != 0:
            r = subprocess.run(["tesseract", str(arq), "stdout", "--psm", "6"], capture_output=True, text=True, timeout=120)
    return r.stdout


def ler_print(imagem: bytes, **extras: Any) -> Carteira:
    texto = ocr(imagem)
    if len(texto.strip()) < 10:
        raise ValueError("Não consegui ler texto no print (imagem pequena ou desfocada?).")
    c = ler_texto(texto, **extras)
    c.avisos.append("Lida de um print por OCR: confira os valores (OCR pode errar dígitos).")
    return c


# ---------------------------------------------------------------- valores do dia
def avaliar(carteira: Carteira) -> Carteira:
    """Preenche o valor de posições com quantidade × cotação do dia (ações, FIIs, ETFs)."""
    from quiron.servicos.mercado import cotacoes

    for p in carteira.posicoes:
        if p.valor or not p.quantidade or not p.ticker:
            continue
        try:
            cot = cotacoes.cotacao(p.ticker)
            p.valor = round(cot.preco * p.quantidade, 2)
        except Exception as e:  # noqa: BLE001
            carteira.avisos.append(f"{p.ticker}: sem cotação agora ({type(e).__name__}); posição sem valor.")
    sem_valor = [p.nome for p in carteira.posicoes if not p.valor]
    if sem_valor:
        carteira.avisos.append("Sem valor (fora das contas): " + ", ".join(sem_valor))
        carteira.posicoes = [p for p in carteira.posicoes if p.valor]
    return carteira
